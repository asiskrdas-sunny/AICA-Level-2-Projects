"""
Comprehensive Verification and Accuracy Test Suite for Detailed Project Report
Tests both:
1. Exact cell-by-cell numerical match against reference model:
   'Hotel Case Study Financial Model Aug 2026.xlsx'
2. Full mathematical and accounting audit of the application's engine:
   - Balance Sheet equation (Assets == Liabilities & Net Worth)
   - Cash Flow statement cash reconciliation
   - Loan repayment amortization to zero
   - NPV at calculated IRR == 0
   - Break-even equations
   - Excel export generation and structure
"""

import os
import sys
import copy
import math
import openpyxl

import DetailedProjectReport as dpr

def run_comprehensive_tests():
    excel_path = r'c:\Users\91943\OneDrive\Desktop\Caps Stone Project\Hotel Case Study Financial Model Aug 2026.xlsx'
    if not os.path.exists(excel_path):
        print(f"Error: Reference file not found at {excel_path}")
        return False
        
    temp_ref = None
    try:
        wb = openpyxl.load_workbook(excel_path, data_only=True)
    except Exception:
        import ctypes, tempfile
        temp_ref = os.path.join(tempfile.gettempdir(), "_Hotel_Ref_Temp.xlsx")
        ctypes.windll.kernel32.CopyFileW(excel_path, temp_ref, False)
        wb = openpyxl.load_workbook(temp_ref, data_only=True)
    
    total_checks = 0
    passed_checks = 0
    discrepancies = []

    def check(name, computed, expected, tol=0.01):
        nonlocal total_checks, passed_checks
        total_checks += 1
        if expected is None:
            if computed is None:
                passed_checks += 1
                return True
            else:
                discrepancies.append(f"{name}: computed {computed}, expected None")
                return False
        if computed is None:
            discrepancies.append(f"{name}: computed None, expected {expected}")
            return False
        
        diff = abs(computed - expected)
        if diff <= tol:
            passed_checks += 1
            return True
        else:
            pct_err = (diff / abs(expected)) * 100 if abs(expected) > 1e-6 else 0
            discrepancies.append(f"{name}: computed {computed:.4f}, expected {expected:.4f} (diff: {diff:.4f}, err: {pct_err:.2f}%)")
            return False

    print("=" * 80)
    print(" DETAILED PROJECT REPORT - AUDIT & VERIFICATION SUITE")
    print("=" * 80)

    # =========================================================================
    # PART 1: EXACT NUMERICAL VERIFICATION AGAINST REFERENCE MODEL
    # =========================================================================
    print("\n--- PART 1: VERIFICATION AGAINST 'Hotel Case Study Financial Model Aug 2026.xlsx' ---")
    
    ws_assump = wb['Assumption Sheet']
    ws_is = wb['Project Income Statement']
    ws_bs = wb['Projected Balance Sheet']
    ws_cf = wb['Projected Cash FLow Statement']

    # Reference stepped repayment schedule from Excel: [2.5%, 5%, 7.5%, 10%, 12.5% x 6]
    excel_rep_pct = [ws_is.cell(34, c).value for c in range(2, 12)]

    a1 = copy.deepcopy(dpr.DEFAULTS)
    a1['repay_type'] = 'stepped'
    R1 = dpr.compute(a1)

    # 1.1 Project Cost & Financing
    print("1.1 Checking Project Cost & Means of Finance...")
    check("Core Cost", R1['core'], ws_assump.cell(9, 2).value)
    check("Core Promoters", R1['core_prom'], ws_assump.cell(9, 4).value)
    check("Core Bank Loan", R1['core_bank'], ws_assump.cell(9, 5).value)
    check("IDC Total", R1['idc'], ws_assump.cell(10, 2).value)
    check("Total Project Cost", R1['total_cost'], ws_assump.cell(11, 2).value)
    check("Total Promoters Contribution", R1['total_prom'], ws_assump.cell(11, 4).value)
    check("Total Bank Loan", R1['total_bank'], ws_assump.cell(11, 5).value)

    # 1.2 Income Statement (10 Operating Years)
    print("1.2 Checking Income Statement (Revenues, Expenses, PBT, Tax, PAT)...")
    for i in range(10):
        col = 2 + i
        yr = ws_is.cell(2, col).value
        # Occupancy & Tariff
        check(f"Occupancy ({yr})", R1['occ'][i], ws_is.cell(4, col).value, tol=0.001)
        check(f"Tariff ({yr})", R1['tariff'][i], ws_is.cell(5, col).value, tol=0.5)
        # Revenues
        check(f"Room Rev ({yr})", R1['room'][i], ws_is.cell(7, col).value)
        check(f"F&B Rev ({yr})", R1['fb'][i], ws_is.cell(8, col).value)
        check(f"Other Rev ({yr})", R1['oth'][i], ws_is.cell(9, col).value)
        check(f"Total Rev ({yr})", R1['rev'][i], ws_is.cell(10, col).value)
        # Variable Expenses
        check(f"Var F&B ({yr})", R1['v_fb'][i], ws_is.cell(13, col).value)
        check(f"Var Emp ({yr})", R1['v_emp'][i], ws_is.cell(14, col).value)
        check(f"Var Power ({yr})", R1['v_pow'][i], ws_is.cell(15, col).value)
        check(f"Var Other ({yr})", R1['v_oth'][i], ws_is.cell(16, col).value)
        check(f"Var Service ({yr})", R1['v_srv'][i], ws_is.cell(17, col).value)
        check(f"Total Var Exp ({yr})", R1['var_tot'][i], ws_is.cell(18, col).value)
        # Fixed Expenses
        check(f"Admin Exp ({yr})", R1['admin'][i], ws_is.cell(20, col).value)
        check(f"Depreciation ({yr})", R1['depr'][i], ws_is.cell(21, col).value)
        check(f"Interest ({yr})", R1['l_int'][i], ws_is.cell(22, col).value)
        check(f"Total Fixed Exp ({yr})", R1['fixed'][i], ws_is.cell(23, col).value)
        # PBT, Tax, PAT
        check(f"Total Exp ({yr})", R1['total_exp'][i], ws_is.cell(24, col).value)
        check(f"PBT ({yr})", R1['pbt'][i], ws_is.cell(25, col).value)
        check(f"Tax ({yr})", R1['tax'][i], ws_is.cell(26, col).value)
        check(f"PAT ({yr})", R1['pat'][i], ws_is.cell(27, col).value)

    # 1.3 Term Loan Schedule, DSCR, IRR, Payback & Break-Even
    print("1.3 Checking Term Loan, DSCR, IRR, Payback & Break-Even...")
    for i in range(10):
        col = 2 + i
        yr = ws_is.cell(2, col).value
        check(f"Loan Open ({yr})", R1['l_open'][i], ws_is.cell(35, col).value)
        check(f"Loan Repay ({yr})", R1['l_rep'][i], ws_is.cell(36, col).value)
        check(f"Loan Close ({yr})", R1['l_close'][i], ws_is.cell(37, col).value)
        check(f"Loan Avg ({yr})", R1['l_avg'][i], ws_is.cell(38, col).value)
        check(f"Loan Int ({yr})", R1['l_int'][i], ws_is.cell(39, col).value)
        check(f"CFADS ({yr})", R1['cfads'][i], ws_is.cell(47, col).value)
        check(f"Debt Repay Total ({yr})", R1['debt_service'][i], ws_is.cell(50, col).value)
        check(f"DSCR ({yr})", R1['dscr'][i], ws_is.cell(51, col).value, tol=0.01)
        check(f"Contribution ({yr})", R1['contrib'][i], ws_is.cell(77, col).value)
        check(f"P/V Ratio ({yr})", R1['pv'][i], ws_is.cell(78, col).value, tol=0.001)
        check(f"BEP ({yr})", R1['bep'][i], ws_is.cell(81, col).value)
        check(f"BEP % ({yr})", R1['bep_pct'][i], ws_is.cell(82, col).value, tol=0.005)
        check(f"Cash BEP ({yr})", R1['cbep'][i], ws_is.cell(84, col).value)
        check(f"Cash BEP % ({yr})", R1['cbep_pct'][i], ws_is.cell(85, col).value, tol=0.005)

    check("Average DSCR", R1['avg_dscr'], ws_is.cell(52, 2).value, tol=0.01)
    check("Project IRR", R1['irr'], ws_is.cell(62, 2).value, tol=0.001)
    check("Payback Period", R1['payback'], ws_is.cell(70, 2).value, tol=0.05)

    # 1.4 Projected Balance Sheet (12 Years: 2 Const + 10 Ops)
    print("1.4 Checking Projected Balance Sheet...")
    for i in range(12):
        col = 2 + i
        yr = ws_bs.cell(2, col).value
        check(f"Gross Block ({yr})", R1['gross'][i], ws_bs.cell(5, col).value)
        check(f"Acc Dep ({yr})", R1['accdep'][i], ws_bs.cell(6, col).value)
        check(f"Net Block ({yr})", R1['net'][i], ws_bs.cell(7, col).value)
        check(f"CWIP ({yr})", R1['cwip'][i], ws_bs.cell(8, col).value)
        check(f"Cash & Bank ({yr})", R1['cash'][i], ws_bs.cell(9, col).value)
        check(f"Total Assets ({yr})", R1['assets'][i], ws_bs.cell(10, col).value)
        check(f"Equity ({yr})", R1['equity'][i], ws_bs.cell(11, col).value)
        check(f"Reserves ({yr})", R1['reserves'][i], ws_bs.cell(12, col).value)
        check(f"Net Worth ({yr})", R1['networth'][i], ws_bs.cell(13, col).value)
        check(f"Bank Loan ({yr})", R1['bank_loan'][i], ws_bs.cell(14, col).value)
        check(f"Total Liab ({yr})", R1['liab'][i], ws_bs.cell(15, col).value)
        check(f"BS Difference ({yr})", R1['diff'][i], ws_bs.cell(16, col).value or 0.0, tol=1e-5)

    # 1.5 Projected Cash Flow Statement (12 Years)
    print("1.5 Checking Projected Cash Flow Statement...")
    for i in range(12):
        col = 2 + i
        yr = ws_cf.cell(2, col).value
        check(f"CF PAT ({yr})", R1['cf_npat'][i], ws_cf.cell(5, col).value)
        check(f"CF Dep ({yr})", R1['cf_dep'][i], ws_cf.cell(6, col).value)
        check(f"CF Equity ({yr})", R1['cf_cap'][i], ws_cf.cell(7, col).value)
        check(f"CF Debt ({yr})", R1['cf_debt'][i], ws_cf.cell(8, col).value or 0.0)
        check(f"CF Total Funds ({yr})", R1['cf_funds'][i], ws_cf.cell(9, col).value)
        check(f"CF Capex ({yr})", R1['cf_fa'][i], ws_cf.cell(11, col).value)
        check(f"CF Repayment ({yr})", R1['cf_rep'][i], ws_cf.cell(12, col).value)
        check(f"CF Total Out ({yr})", R1['cf_out'][i], ws_cf.cell(13, col).value)
        check(f"CF Surplus ({yr})", R1['cf_surplus'][i], ws_cf.cell(14, col).value)
        check(f"CF Opening ({yr})", R1['cf_open'][i], ws_cf.cell(15, col).value)
        check(f"CF Closing ({yr})", R1['cf_close'][i], ws_cf.cell(16, col).value)


    # =========================================================================
    # PART 2: ACCOUNTING INVARIANTS & MATHEMATICAL INTEGRITY (DEFAULT ENGINE)
    # =========================================================================
    print("\n--- PART 2: INTERNAL ACCOUNTING & MATHEMATICAL INVARIANTS (DEFAULT ENGINE) ---")
    a2 = copy.deepcopy(dpr.DEFAULTS)
    a2['repay_type'] = 'equal'
    R2 = dpr.compute(a2)

    # 2.1 Balance Sheet Balance: Total Assets == Total Liabilities (Every Year)
    print("2.1 Verifying Balance Sheet Balance: Assets == Liabilities...")
    for i, d in enumerate(R2['diff']):
        check(f"BS Identity Year {i+1} (Difference == 0.00)", d, 0.0, tol=1e-6)

    # 2.2 Cash Flow Reconciliation: Closing Cash == Balance Sheet Cash
    print("2.2 Verifying Cash Flow Statement Closing Cash == Balance Sheet Cash...")
    for i in range(12):
        check(f"Cash Reconciliation Year {i+1}", R2['cash'][i], R2['cf_close'][i], tol=1e-6)

    # 2.3 Financing Invariant: Total Project Cost == Promoters Contribution + Bank Loan
    print("2.3 Verifying Total Project Cost == Promoters + Bank Loan...")
    check("Financing Balance Identity", R2['total_cost'], R2['total_prom'] + R2['total_bank'], tol=1e-6)

    # 2.4 Loan Repayment Invariant: Total Repaid == Total Borrowed & Closing Loan == 0
    print("2.4 Verifying Loan Repayment Invariants...")
    check("Total Loan Repaid == Total Loan Borrowed", sum(R2['l_rep']), R2['loan'], tol=1e-6)
    check("Final Year Closing Loan Balance == 0", R2['l_close'][-1], 0.0, tol=1e-6)

    # 2.5 Fixed Asset & Depreciation Invariants
    print("2.5 Verifying Fixed Asset Accounting Identities...")
    for i in range(12):
        check(f"Net Block == Gross Block - AccDep Year {i+1}", R2['net'][i], R2['gross'][i] - R2['accdep'][i], tol=1e-6)
    for yr_idx in range(2, 12):
        check(f"CWIP == 0 in Operating Year {yr_idx-1}", R2['cwip'][yr_idx], 0.0, tol=1e-6)

    # 2.6 Net Worth Invariant: Net Worth == Equity + Accumulated Reserves
    print("2.6 Verifying Net Worth Identity...")
    for i in range(12):
        check(f"Net Worth == Equity + Reserves Year {i+1}", R2['networth'][i], R2['equity'][i] + R2['reserves'][i], tol=1e-6)

    # 2.7 IRR Verification: NPV at calculated IRR should be exactly 0
    print("2.7 Verifying IRR Equation: NPV(IRR) == 0...")
    irr_val = R2['irr']
    npv_at_irr = sum(cf / ((1.0 + irr_val) ** t) for t, cf in enumerate(R2['irr_tot']))
    check("NPV at Project IRR is zero", npv_at_irr, 0.0, tol=1e-5)

    # 2.8 DSCR Calculation Invariant: DSCR == CFADS / Debt Service
    print("2.8 Verifying DSCR Identity...")
    for i in range(10):
        expected_dscr = R2['cfads'][i] / R2['debt_service'][i]
        check(f"DSCR Formula Year {i+1}", R2['dscr'][i], expected_dscr, tol=1e-9)

    # 2.9 Break-Even Analysis Invariant
    print("2.9 Verifying Break-Even Equations...")
    for i in range(10):
        expected_bep = R2['fixed'][i] / R2['pv'][i]
        check(f"BEP Formula Year {i+1}", R2['bep'][i], expected_bep, tol=1e-9)
        expected_cbep = R2['fixed_cash'][i] / R2['pv'][i]
        check(f"Cash BEP Formula Year {i+1}", R2['cbep'][i], expected_cbep, tol=1e-9)


    # =========================================================================
    # PART 3: REPORTS & EXCEL EXPORT INTEGRITY
    # =========================================================================
    print("\n--- PART 3: REPORT BUILDER & EXCEL EXPORT INTEGRITY ---")
    IS, BS, CF = dpr.build_reports(a2, R2)
    check("Income Statement rows exist", len(IS) > 0, True)
    check("Balance Sheet rows exist", len(BS) > 0, True)
    check("Cash Flow Statement rows exist", len(CF) > 0, True)
    
    # Check Assumption rows for export
    assump_rows = dpr.assumption_rows(a2, R2)
    check("Assumption summary rows exist", len(assump_rows) > 0, True)

    # Test Excel Export Workbook generation (10-sheet live-formula model)
    print("3.1 Testing 10-Sheet Live-Formula Excel Export generation and schema...")
    test_xlsx = "Test_Audit_Export.xlsx"
    dpr.export_to_excel(test_xlsx, a2, R2)
    check("Export file created on disk", os.path.exists(test_xlsx), True)
    
    # Reload and verify
    wb_read = openpyxl.load_workbook(test_xlsx, data_only=False)
    check("Exported workbook has exactly 10 sheets", len(wb_read.sheetnames), 10)
    expected_10_sheets = [
        "Cover", "Assumptions", "Project Cost & Finance", "Depreciation",
        "Loan Repayment Schedule", "Projected P&L", "Projected Balance Sheet",
        "Cash Flow", "Ratios & DSCR", "Break-even"
    ]
    for expected_name in expected_10_sheets:
        check(f"Sheet '{expected_name}' in export", expected_name in wb_read.sheetnames, True)
    
    # Audit formulas across sheets
    total_formulas = 0
    ref_errors = 0
    for sname in wb_read.sheetnames:
        ws_test = wb_read[sname]
        for row in ws_test.iter_rows(values_only=False):
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    if cell.value.startswith('='):
                        total_formulas += 1
                    if '#REF!' in cell.value or '#NAME?' in cell.value:
                        ref_errors += 1
    check("Total live formulas > 1000", total_formulas > 1000, True)
    check("Zero formula reference errors (#REF! / #NAME?)", ref_errors, 0)
    
    # Check input styling on Assumptions sheet
    ws_assump = wb_read["Assumptions"]
    input_styled = 0
    for r in range(4, ws_assump.max_row + 1):
        for c in range(2, ws_assump.max_column + 1):
            cell = ws_assump.cell(r, c)
            if cell.fill and cell.fill.start_color and cell.fill.start_color.rgb == '00FFF2CC':
                if cell.font and cell.font.color and cell.font.color.rgb == '00002060':
                    input_styled += 1
    check("Assumptions input cells styled (Yellow fill + Blue font)", input_styled > 50, True)
    
    # Check Balance Sheet check cell formula
    ws_bs = wb_read["Projected Balance Sheet"]
    diff_val = ws_bs.cell(19, 3).value
    check("Balance Sheet check formula is difference (Assets - Liab)", isinstance(diff_val, str) and diff_val.startswith('='), True)
    
    # Check charts on Ratios & DSCR sheet
    ws_ratios = wb_read["Ratios & DSCR"]
    check("Ratios & DSCR sheet contains live openpyxl charts", len(ws_ratios._charts), 2)
    
    if os.path.exists(test_xlsx):
        os.remove(test_xlsx)

    print("3.2 Testing PDF Export generation and structure...")
    test_pdf = "Test_Audit_Export.pdf"
    dpr.export_to_pdf(test_pdf, a2, R2)
    check("PDF export file created on disk", os.path.exists(test_pdf), True)
    check("PDF file size > 10 KB", os.path.getsize(test_pdf) > 10000, True)

    try:
        import pymupdf as fitz
        pdf_doc = fitz.open(test_pdf)
        check("PDF has exactly 7 executive pages", len(pdf_doc), 7)
        check("Page 1 contains DPR Title", "DETAILED PROJECT REPORT" in pdf_doc[0].get_text(), True)
        check("Page 1 contains Entity Name", a2.get("entity_name", "Grand Horizon") in pdf_doc[0].get_text(), True)
        check("Page 1 contains Bank Name", a2.get("bank_name", "State Bank of India") in pdf_doc[0].get_text(), True)
        check("Page 2 contains Executive Dashboard", "EXECUTIVE DASHBOARD" in pdf_doc[1].get_text() or "Key Financial" in pdf_doc[1].get_text(), True)
        check("Page 3 contains Assumptions Title", "PROJECT ASSUMPTIONS" in pdf_doc[2].get_text(), True)
        check("Page 4 contains Income Statement", "PROJECT INCOME STATEMENT" in pdf_doc[3].get_text(), True)
        check("Page 5 contains Debt Service Schedule", "TERM LOAN REPAYMENT" in pdf_doc[4].get_text(), True)
        check("Page 6 contains Profitability & Viability", "PROJECT PROFITABILITY" in pdf_doc[5].get_text(), True)
        check("Page 7 contains Balance Sheet", "PROJECTED BALANCE SHEET" in pdf_doc[6].get_text(), True)
        pdf_doc.close()
    except Exception as e:
        check(f"PDF content verification ({e})", False, True)

    if os.path.exists(test_pdf):
        os.remove(test_pdf)
    if temp_ref and os.path.exists(temp_ref):
        try:
            os.remove(temp_ref)
        except Exception:
            pass

    print("\n" + "=" * 80)
    print(f" FINAL TEST REPORT: {total_checks} CHECKS EXECUTED")
    print(f" PASSED: {passed_checks} / {total_checks} ({passed_checks / total_checks * 100:.1f}%)")
    if discrepancies:
        print(f" DISCREPANCIES ({len(discrepancies)}):")
        for d in discrepancies[:25]:
            print("   *", d)
        if len(discrepancies) > 25:
            print(f"   ... and {len(discrepancies) - 25} more.")
        return False
    else:
        print(" ALL CHECKS PASSED WITH 100% ACCURACY!")
        print(" Output is mathematically sound, accounting-compliant, and perfectly matched.")
        print("=" * 80)
        return True

if __name__ == '__main__':
    success = run_comprehensive_tests()
    sys.exit(0 if success else 1)
